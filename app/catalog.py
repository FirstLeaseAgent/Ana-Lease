"""Private, versioned capture catalogues. Existing requests keep a snapshot."""
import copy
import re
from io import BytesIO
from .rules import FIELDS
from . import shareholders

ROLES={'solicitante','contacto','aval','representante','accionista'}
DOC_ROLES={'Solicitante','Aval','RepresentanteLegal','Accionista'}
TYPES={'text','email','phone','url','select'}
CODE=re.compile(r'^[a-z][a-z_]{0,79}$')

def baseline():
    from .conversation import ORDER,QUESTIONS
    from .documents import CATALOG
    rows=[]
    for role,subjects in FIELDS.items():
        for subject,fields in subjects.items():
            for i,code in enumerate(ORDER[role]):
                if code in fields:
                    rows.append({'role':role,'subject_type':subject,'code':code,'label':code.replace('_',' '),'question':QUESTIONS[code],'type':{'correo_contacto':'email','telefono':'phone','pagina_web':'url'}.get(code,'text'),'order':i+1,'enabled':True,'options':[],'reuse_from':'razon_social' if code=='nombre_comercial' else ''})
    return {'fields':rows,'documents':copy.deepcopy(CATALOG)}

def active(conn):
    row=conn.execute('SELECT id,body FROM capture_catalogs ORDER BY id DESC LIMIT 1').fetchone()
    if row and 'body' in row:
        config=copy.deepcopy(row['body'])
        if not any(r['role']=='accionista' for r in config['fields']):
            config['fields'].extend(r for r in baseline()['fields'] if r['role']=='accionista')
        return row['id'],config
    return 0,baseline()

def snapshot(conn,intake_id):
    row=conn.execute('SELECT catalog_snapshot FROM intakes WHERE id=%s',(intake_id,)).fetchone()
    return row.get('catalog_snapshot') or baseline() if row else baseline()

def field_rows(person,subject=None):
    config=person.get('catalog')
    if config is None:return None
    rows=config['fields']
    # Existing intake snapshots acquire defaults only for an explicitly added new role.
    if person['role']=='accionista' and not any(r['role']=='accionista' for r in rows):
        rows=baseline()['fields']
    return sorted([r for r in rows if r['role']==person['role'] and r['subject_type']==(subject or person['subject_type']) and r['enabled']],key=lambda r:(r['order'],r['code']))

def fields_for(person,subject=None):
    rows=field_rows(person,subject)
    return {r['code'] for r in rows} if rows is not None else set(FIELDS[person['role']][subject or person['subject_type']])

def question_for(person,code):
    from .conversation import QUESTIONS
    rows=field_rows(person)
    row=next((r for r in rows or [] if r['code']==code),None)
    text=row['question'] if row else QUESTIONS.get(code,'¿Cuál es '+code.replace('_',' ')+'?')
    if row and row['options']:text+=' Opciones: '+', '.join(row['options'])+'.'
    return text

def reuse_source(person,code):
    rows=field_rows(person)
    return next((r['reuse_from'] for r in rows or [] if r['code']==code),'') if rows is not None else ('razon_social' if code=='nombre_comercial' else '')

def validate_value(person,code,value):
    from .pending import unavailable
    if unavailable(value,code):raise ValueError('Ese dato no se guardó. Escribe «no lo tengo» para dejarlo pendiente o pide que te contacten.')
    if code=='curp' and len(value.strip())!=18:
        raise ValueError(f'La CURP debe tener 18 caracteres; escribiste {len(value.strip())}. Rectifícala o escribe «no lo tengo» para continuar. También puedes pedir que te contacten.')
    if person['role']=='accionista':
        if code=='porcentaje_participacion':shareholders.percentage(value)
        if code=='curp' and not shareholders.CURP.fullmatch(value.upper()):
            raise ValueError('Revisa el formato de la CURP (18 caracteres)')
    row=next((r for r in field_rows(person) or [] if r['code']==code),None)
    if row and row['type']=='select' and value.casefold() not in {s.casefold() for s in row['options']}:
        raise ValueError('Selecciona una de las opciones permitidas')

def validate(config):
    if not isinstance(config,dict) or set(config)!={'fields','documents'}:raise ValueError('El catálogo debe contener fields y documents')
    result={'fields':[],'documents':[]}
    if not isinstance(config['fields'],list) or not 1<=len(config['fields'])<=320:raise ValueError('Se permiten entre 1 y 320 preguntas')
    if not isinstance(config['documents'],list) or not 1<=len(config['documents'])<=150:raise ValueError('Se permiten entre 1 y 150 documentos')
    seen=set()
    def text(value,key,limit=500,empty=False):
        if not isinstance(value,str) or len(value)>limit or (not empty and not value.strip()) or value.startswith(('=','+','@')):raise ValueError('Revisa '+key)
        return value.strip()
    def boolean(value,key):
        if type(value) is not bool:raise ValueError(key+' debe ser verdadero o falso')
        return value
    def order(value):
        if type(value) is not int or not 1<=value<=1000:raise ValueError('El orden debe ser un entero de 1 a 1000')
        return value
    for raw in config['fields']:
        if not isinstance(raw,dict):raise ValueError('Pregunta inválida')
        role=raw.get('role');subject=raw.get('subject_type');code=raw.get('code')
        if role not in ROLES or subject not in {'PF','PM'} or (role in {'contacto','representante'} and subject!='PF'):raise ValueError('Rol o tipo de persona inválido')
        if not isinstance(code,str) or not CODE.fullmatch(code) or code=='rfc':raise ValueError('Código de campo inválido; RFC es un campo protegido del sistema')
        key=(role,subject,code)
        if key in seen:raise ValueError('Pregunta duplicada: '+code)
        seen.add(key)
        kind=raw.get('type','text')
        if kind not in TYPES:raise ValueError('Tipo de campo inválido')
        options=raw.get('options',[])
        if not isinstance(options,list) or len(options)>30:raise ValueError('Opciones inválidas')
        options=[text(o,'opción',100) for o in options]
        if kind=='select' and not options:raise ValueError('Una lista necesita opciones')
        source=raw.get('reuse_from','') or ''
        if source and (not isinstance(source,str) or not CODE.fullmatch(source) or source==code):raise ValueError('Origen de reutilización inválido')
        result['fields'].append({'role':role,'subject_type':subject,'code':code,'label':text(raw.get('label'),'nombre',120),'question':text(raw.get('question'),'pregunta'),'type':kind,'order':order(raw.get('order')),'enabled':boolean(raw.get('enabled',True),'enabled'),'options':options,'reuse_from':source})
    for subject,required in [('PF',{'nombre','curp','porcentaje_participacion'}),('PM',{'razon_social','porcentaje_participacion'})]:
        configured=[r for r in result['fields'] if r['role']=='accionista' and r['subject_type']==subject]
        if any(r['role']=='accionista' for r in result['fields']) and not required <= {r['code'] for r in configured if r['enabled']}:
            raise ValueError('Accionista '+subject+' requiere identidad y porcentaje de participación activos')
    for row in result['fields']:
        if row['enabled'] and row['reuse_from'] and not any(r['role']==row['role'] and r['subject_type']==row['subject_type'] and r['code']==row['reuse_from'] and r['enabled'] and r['type']==row['type'] for r in result['fields']):raise ValueError('El origen de reutilización debe ser un campo activo del mismo rol, tipo de persona y tipo de dato')
    seen=set()
    for raw in config['documents']:
        if not isinstance(raw,dict):raise ValueError('Documento inválido')
        scope=raw.get('scope');role=raw.get('role');code=raw.get('code')
        if scope not in {'PF','PM'} or role not in DOC_ROLES or not isinstance(code,str) or not CODE.fullmatch(code):raise ValueError('Rol, alcance o código documental inválido')
        key=(scope,role,code)
        if key in seen:raise ValueError('Documento duplicado: '+code)
        seen.add(key)
        field=raw.get('dependency_field') or None;value=raw.get('dependency_value') or None
        if field and (not isinstance(field,str) or not CODE.fullmatch(field) or not value):raise ValueError('Condición documental inválida')
        if value and not field:raise ValueError('Falta el campo de la condición')
        result['documents'].append({'scope':scope,'role':role,'code':code,'label':text(raw.get('label'),'documento',150),'order':order(raw.get('order')),'required':boolean(raw.get('required'),'required'),'dependency_field':field,'dependency_value':text(value,'valor de condición',100) if value else None})
    return result

FIELD_COLUMNS=['role','subject_type','code','label','question','type','order','enabled','reuse_from','options']
DOC_COLUMNS=['scope','role','code','label','order','required','dependency_field','dependency_value']

def workbook(config):
    from openpyxl import Workbook
    book=Workbook();book.remove(book.active)
    for name,key,columns in [('Preguntas','fields',FIELD_COLUMNS),('Documentos','documents',DOC_COLUMNS)]:
        sheet=book.create_sheet(name);sheet.append(columns);sheet.freeze_panes='A2'
        for row in config[key]:sheet.append([' | '.join(row[c]) if c=='options' else row.get(c) for c in columns])
        sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in col)+2))
    out=BytesIO();book.save(out);return out.getvalue()

def parse_workbook(content):
    import zipfile
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist())>100 or sum(f.file_size for f in archive.infolist())>4*1024*1024:raise ValueError('El Excel descomprimido excede el límite')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
        result={}
        for name,key,columns in [('Preguntas','fields',FIELD_COLUMNS),('Documentos','documents',DOC_COLUMNS)]:
            if name not in book.sheetnames:raise ValueError('Falta la hoja '+name)
            sheet=book[name]
            if sheet.max_row>500 or sheet.max_column>20:raise ValueError('La hoja es demasiado grande')
            rows=sheet.iter_rows();header=[cell.value for cell in next(rows)]
            if header!=columns:raise ValueError('Los encabezados de '+name+' no coinciden con la plantilla')
            result[key]=[]
            for cells in rows:
                if any(c.data_type=='f' for c in cells):raise ValueError('La plantilla no admite fórmulas')
                values=[c.value for c in cells]
                if all(v is None for v in values):continue
                row=dict(zip(columns,values))
                for c in ['enabled','required']:
                    if c in row:
                        v=row[c]
                        if isinstance(v,str) and v.strip().casefold() in {'si','sí','true','1','no','false','0'}:v=v.strip().casefold() in {'si','sí','true','1'}
                        if type(v) is not bool:raise ValueError(c+' debe ser TRUE/FALSE o Sí/No')
                        row[c]=v
                if key=='fields':
                    row['options']=[s.strip() for s in str(row.get('options') or '').split('|') if s.strip()];row['reuse_from']=row.get('reuse_from') or ''
                result[key].append(row)
        book.close()
        return validate(result)
    except ValueError:raise
    except Exception:raise ValueError('No pudimos leer el Excel. Usa la plantilla .xlsx sin macros.')
